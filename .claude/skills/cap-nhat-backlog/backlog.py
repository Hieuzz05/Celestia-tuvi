"""Ghi PRODUCT-BACKLOG.xlsx bằng openpyxl — dùng cho skill /cap-nhat-backlog.

Đầu vào là MỘT tệp JSON (viết bằng công cụ Write, đừng gõ JSON trên dòng lệnh:
harness nuốt một lớp dấu gạch chéo, xem docs/bay/moi-truong.md).

    python .claude/skills/cap-nhat-backlog/backlog.py xem               # ID kế tiếp, 3 dòng cuối mỗi sheet
    python .claude/skills/cap-nhat-backlog/backlog.py tim CEL-186       # in dòng Backlog + các bước Logic của một ID
    python .claude/skills/cap-nhat-backlog/backlog.py ghi <tep.json>    # áp thay đổi

Dạng tệp JSON (mọi khoá đều tuỳ chọn, trừ nhat_ky):
{
  "backlog_sua": [{"id": "CEL-186", "cot": {"Trạng thái": "Đang chạy", "Việc còn lại / rủi ro": "..."}}],
  "backlog_moi": [{"Nhóm": "...", "Tính năng": "...", ...}],      # ID tự cấp nếu không ghi
  "logic_them": [{"Mã": "CEL-186", "Luồng": "...", "#": "1", "Bước": "...", "Logic cụ thể": "...",
                  "Luật bất biến / bẫy đã trả giá": "...", "Tệp": "..."}],
  "logic_sua":  [{"Mã": "CEL-186", "#": "1", "cot": {"Logic cụ thể": "..."}}],
  "nhat_ky": {"Commit": "abc1234", "ID liên quan": "CEL-186", "Thay đổi gì": "...", "Vì sao": "...",
              "Ai làm": "Claude (máy 1)"},
  "commit": "abc1234"          # điền cột Commit của mọi dòng Backlog vừa sửa/thêm
}
Cột "Cập nhật" của Backlog và "Ngày" của Nhật ký tự điền ngày hôm nay (dd/mm/yyyy).
"""
import json
import re
import subprocess
import sys
from datetime import date

import openpyxl

TEP = 'PRODUCT-BACKLOG.xlsx'
HOM_NAY = date.today().strftime('%d/%m/%Y')

try:
    sys.stdout.reconfigure(encoding='utf-8')  # cp1252 của Windows làm sập khi in tiếng Việt
except AttributeError:
    pass


def tieu_de(ws):
    return [c.value for c in ws[1]]


def cot(ws, ten):
    td = tieu_de(ws)
    if ten not in td:
        sys.exit(f'Sheet "{ws.title}" không có cột "{ten}". Có: {td}')
    return td.index(ten) + 1


def _so_trong_sheet(ws):
    return [int(m.group(1)) for (v,) in ws.iter_rows(min_row=2, max_col=1, values_only=True)
            if v and (m := re.fullmatch(r'CEL-(\d+)', str(v).strip()))]


def _so_tren_nhanh_remote():
    """ID trong Backlog của mọi nhánh remote (sau git fetch). Nhánh kia có thể đã push ID trong
    backlog mà chưa có tiêu đề commit nào ghi ID đó, nên chỉ quét git log là không đủ."""
    import io
    so, da_doc = [], set()
    try:
        subprocess.run(['git', 'fetch', '--quiet', 'origin'], capture_output=True, timeout=30)
    except (OSError, subprocess.TimeoutExpired):
        print('(không fetch được origin — ID kế tiếp có thể đã bị nhánh khác giữ)')
    try:
        refs = subprocess.run(['git', 'for-each-ref', '--format=%(refname)', 'refs/remotes/origin'],
                              capture_output=True, text=True).stdout.split()
    except OSError:
        return so
    for ref in refs:
        blob = subprocess.run(['git', 'rev-parse', '--verify', '--quiet', f'{ref}:{TEP}'],
                              capture_output=True, text=True).stdout.strip()
        if not blob or blob in da_doc:
            continue
        da_doc.add(blob)
        du = subprocess.run(['git', 'cat-file', 'blob', blob], capture_output=True).stdout
        try:
            wb = openpyxl.load_workbook(io.BytesIO(du), read_only=True)
            so += _so_trong_sheet(wb['Backlog'])
        except Exception:  # bản xlsx cũ hỏng / thiếu sheet: bỏ qua, không chặn việc cấp
            pass
    return so


def id_ke_tiep(ws):
    """Lớn nhất trong sheet, trong Backlog của mọi nhánh remote, VÀ trong commit của mọi nhánh:
    nhánh chưa gộp có thể đã giữ một ID mà bản backlog ở nhánh này chưa thấy (trùng CEL-150, gỡ ở
    87f5b0c; trùng CEL-194, đổi sang CEL-196 ở fb2b023). Nhánh CHƯA push thì không công cụ nào
    thấy được: push nhánh ngay sau commit đầu mang ID mới; CI kiem-id-cel.yml bắt phần còn lọt."""
    so = _so_trong_sheet(ws) + _so_tren_nhanh_remote()
    try:
        log = subprocess.run(['git', 'log', '--all', '--format=%s', '-n', '500'],
                             capture_output=True, text=True, encoding='utf-8').stdout
        so += [int(x) for x in re.findall(r'CEL-(\d+)', log)]
    except OSError:
        pass
    return f'CEL-{max(so) + 1:03d}'


def dong_cua_id(ws, ma):
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, 1).value or '').strip() == ma:
            return r
    return None


def them_dong(ws, du_lieu):
    td = tieu_de(ws)
    la = set(du_lieu) - set(td)
    if la:
        sys.exit(f'Sheet "{ws.title}" không có cột {sorted(la)}')
    r = ws.max_row + 1
    for k, v in du_lieu.items():
        ws.cell(r, td.index(k) + 1, v)
    return r


def xem(wb):
    print('ID kế tiếp:', id_ke_tiep(wb['Backlog']))
    print('  → dùng ID này thì push nhánh ngay sau commit đầu mang nó: nhánh chưa push thì máy kia không thấy.')
    for ten in ('Backlog', 'Logic chi tiết', 'Nhật ký thay đổi'):
        ws = wb[ten]
        print(f'\n== {ten} ({ws.max_row - 1} dòng)')
        for row in ws.iter_rows(min_row=max(2, ws.max_row - 2), values_only=True):
            print('  ', [str(c)[:50] if c is not None else '' for c in row[:5]])


def tim(wb, ma):
    ws = wb['Backlog']
    r = dong_cua_id(ws, ma)
    if r:
        for k, c in zip(tieu_de(ws), ws[r]):
            print(f'{k}: {c.value}')
    else:
        print(f'Backlog: không có {ma}')
    print('\n== Logic chi tiết')
    for row in wb['Logic chi tiết'].iter_rows(min_row=2, values_only=True):
        if str(row[0] or '').strip() == ma:
            print('  ', [str(c)[:70] if c is not None else '' for c in row])


def ghi(wb, duong_dan):
    with open(duong_dan, encoding='utf-8') as f:
        yc = json.load(f)
    if 'nhat_ky' not in yc:
        sys.exit('Thiếu "nhat_ky": mọi lần ghi backlog đều phải có một dòng Nhật ký thay đổi.')
    if not str(yc['nhat_ky'].get('Vì sao', '')).strip():
        sys.exit('Cột "Vì sao" của Nhật ký đang trống — đó là cột quan trọng nhất, viết vào.')

    bl, lg, nk = wb['Backlog'], wb['Logic chi tiết'], wb['Nhật ký thay đổi']
    commit = yc.get('commit')
    da_cham = []

    for muc in yc.get('backlog_sua', []):
        r = dong_cua_id(bl, muc['id'])
        if not r:
            sys.exit(f'Backlog không có {muc["id"]} — tính năng mới thì dùng "backlog_moi".')
        for k, v in muc['cot'].items():
            bl.cell(r, cot(bl, k), v)
        da_cham.append(r)

    for muc in yc.get('backlog_moi', []):
        muc = dict(muc)
        muc.setdefault('ID', id_ke_tiep(bl))
        if dong_cua_id(bl, muc['ID']):
            sys.exit(f'{muc["ID"]} đã có — không dùng lại ID cũ.')
        da_cham.append(them_dong(bl, muc))
        print('Cấp ID:', muc['ID'])

    for r in da_cham:
        bl.cell(r, cot(bl, 'Cập nhật'), HOM_NAY)
        if commit:
            bl.cell(r, cot(bl, 'Commit'), commit)

    for muc in yc.get('logic_sua', []):
        hit = [r for r in range(2, lg.max_row + 1)
               if str(lg.cell(r, 1).value or '').strip() == muc['Mã']
               and str(lg.cell(r, 3).value or '').strip() == str(muc['#'])]
        if len(hit) != 1:
            sys.exit(f'Logic chi tiết: {muc["Mã"]} bước {muc["#"]} khớp {len(hit)} dòng, cần đúng 1.')
        for k, v in muc['cot'].items():
            lg.cell(hit[0], cot(lg, k), v)

    for muc in yc.get('logic_them', []):
        them_dong(lg, muc)

    them_dong(nk, {'Ngày': HOM_NAY, **yc['nhat_ky']})

    wb.save(TEP)
    print(f'Đã ghi: {len(da_cham)} dòng Backlog, '
          f'{len(yc.get("logic_them", [])) + len(yc.get("logic_sua", []))} dòng Logic, 1 dòng Nhật ký.')


def main():
    if len(sys.argv) < 2 or sys.argv[1] not in ('xem', 'tim', 'ghi'):
        sys.exit(__doc__)
    wb = openpyxl.load_workbook(TEP)
    if sys.argv[1] == 'xem':
        xem(wb)
    elif sys.argv[1] == 'tim':
        tim(wb, sys.argv[2])
    else:
        ghi(wb, sys.argv[2])


if __name__ == '__main__':
    main()
